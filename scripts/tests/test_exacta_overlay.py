"""Contract tests for the Exacta fork layer (exa-* and nature-shared/vn)."""
from __future__ import annotations

import importlib.util
import re
import unittest
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[2]
SKILLS = ROOT / "skills"
FROZEN = ROOT / "frozen" / "skills"
CJK_RE = re.compile(r"[一-鿿]")
VN_SKILLS = sorted(p for p in SKILLS.iterdir() if p.is_dir() and p.name.startswith("exa-"))
VN_SHARED = SKILLS / "nature-shared" / "vn"


def vn_text_files():
    for root in [*VN_SKILLS, VN_SHARED]:
        yield from (p for p in root.rglob("*") if p.is_file() and p.suffix in {".md", ".yaml"})


class VietnameseOverlayContractTests(unittest.TestCase):
    def test_expected_skills_present(self):
        names = {p.name for p in VN_SKILLS}
        expected = {
            "exa-tim-tai-lieu", "exa-doc-sau", "exa-dich-thuat",
            "exa-de-cuong", "exa-doc-song-ngu", "exa-phuc-dap",
            "exa-phan-bien", "exa-thong-ke", "exa-viet-bai",
            "exa-ve-hinh", "exa-trinh-chieu", "exa-du-lieu",
        }
        self.assertEqual(names, expected)

    def test_no_chinese_text_in_vn_layer(self):
        offenders = []
        for path in vn_text_files():
            text = path.read_text(encoding="utf-8")
            # The response overlay names the upstream block it replaces; nothing else may use CJK.
            allowed = {"中文核对", "中文", "中文说明", "中文图注", "中文短标题", "公式索引"}
            stripped = text
            for token in sorted(allowed, key=len, reverse=True):
                stripped = stripped.replace(token, "")
            if CJK_RE.search(stripped):
                offenders.append(str(path.relative_to(ROOT)))
        self.assertEqual(offenders, [])

    def test_base_skill_exists_and_is_not_frozen(self):
        for skill in VN_SKILLS:
            manifest = yaml.safe_load((skill / "manifest.yaml").read_text(encoding="utf-8"))
            base = manifest.get("base_skill")
            if base is None:
                continue
            with self.subTest(skill=skill.name):
                self.assertTrue((SKILLS / base / "SKILL.md").is_file(), base)
                self.assertFalse((FROZEN / base).exists(), f"{base} is frozen")
                skill_md = (skill / "SKILL.md").read_text(encoding="utf-8")
                self.assertIn(f"../{base}/SKILL.md", skill_md)

    def test_every_overlay_loads_language_policy(self):
        for skill in VN_SKILLS:
            manifest = yaml.safe_load((skill / "manifest.yaml").read_text(encoding="utf-8"))
            with self.subTest(skill=skill.name):
                self.assertIn("../nature-shared/vn/language-policy.md", manifest["always_load"])

    def test_vn_skills_do_not_route_to_frozen_files(self):
        frozen_refs = ("zh-to-en.md", "chinese-author", "chinese-mode.md", "openrouter")
        for skill in VN_SKILLS:
            manifest_text = (skill / "manifest.yaml").read_text(encoding="utf-8")
            with self.subTest(skill=skill.name):
                for ref in frozen_refs:
                    self.assertNotIn(ref, manifest_text)

    def test_frozen_skills_are_out_of_the_skills_tree(self):
        # Read the registry so this also holds in published snapshots, which omit frozen/skills.
        registry = (ROOT / "frozen" / "FROZEN.md").read_text(encoding="utf-8")
        section = registry.split("## Frozen skills", 1)[1].split("\n## ", 1)[0]
        names = re.findall(r"^\| `(nature-[^`]+)` \|", section, flags=re.MULTILINE)
        self.assertEqual(len(names), 9)
        for name in names:
            with self.subTest(skill=name):
                self.assertFalse((SKILLS / name).exists())

    def test_edition_markers_are_balanced(self):
        builder = ROOT / "scripts" / "exa_build_edition.py"
        if not builder.is_file():
            self.skipTest("edition builder is not shipped in this edition")
        spec = importlib.util.spec_from_file_location("exa_build_edition", builder)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        paths = [p for p in ROOT.rglob("*") if p.suffix in {".md", ".yaml", ".yml"}
                 and ".git" not in p.parts and "frozen" not in p.parts[len(ROOT.parts):][:1]]
        for path in paths:
            text = path.read_text(encoding="utf-8")
            for edition in ("pro", "community"):
                with self.subTest(file=str(path.relative_to(ROOT)), edition=edition):
                    try:
                        module.apply_markers(text, edition, str(path))
                    except SystemExit as exc:
                        self.fail(str(exc))

    def test_docx_export_generic_profile(self):
        try:
            import docx  # noqa: F401
        except ImportError:
            self.skipTest("python-docx not installed")
        import subprocess
        import sys
        import tempfile

        script = VN_SHARED / "scripts" / "md_to_docx.py"
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "in.md"
            out = Path(tmp) / "out.docx"
            src.write_text("# KẾT QUẢ\n\nBảng 3.1. Mẫu (n = 2)\n| A | B |\n|---|---|\n| 1 | 2 |\n\n"
                           "Nhận xét: mẫu [1,2].\n", encoding="utf-8")
            result = subprocess.run([sys.executable, str(script), str(src), str(out), "--profile", "generic"],
                                    capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue(out.is_file())
            document = docx.Document(str(out))
            self.assertEqual(len(document.tables), 1)
            self.assertTrue(any(p.text.startswith("Bảng 3.1.") for p in document.paragraphs))

    def test_rules_files_present(self):
        for rel in ("AGENTS.md", "CLAUDE.md", "frozen/FROZEN.md"):
            with self.subTest(file=rel):
                self.assertTrue((ROOT / rel).is_file())


if __name__ == "__main__":
    unittest.main()
